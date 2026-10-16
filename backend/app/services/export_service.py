"""行程导出服务：把一份行程渲染成 Word（.docx）或 PDF。

【为什么导出的东西必须"像行程单"，而不是"像 JSON"】
    用户在结果页勾选、拖拽、加偏好，最后要的是**能带走的东西** ——
    发给同行的朋友、打印出来塞进包里、或者扔进自己的归档。
    所以导出物里逐日行程、预算、打包清单一件不能少，
    并且必须带上**出处**与**编排来源**：这两样是行迹防幻觉设计的对外证据，
    导出后就看不见界面了，证据必须落在纸面上。

【为什么 Word 和 PDF 共用一份内容模型】
    如果两边各自遍历 plan 组装文本，改一处忘另一处，
    很快就会出现"Word 里有的 PDF 里没有"。所以先抽 `_content()`，
    两个渲染器只负责把同一份结构翻译成各自的排版指令。
    新增字段时只改一处，两边同时生效。

【中文字体的两套兜底】
    Word：直接指定 East Asian 字体（微软雅黑），由 Word 自己渲染。
    PDF ：reportlab 必须显式注册字体。优先用系统 TTF（微软雅黑，字形好看、有粗体），
          取不到就退回内置 CID 字体 STSong-Light —— 它是 reportlab 自带的，
          不依赖任何系统字体文件，保证在裸机上也能出中文而不是一堆黑框。
"""

from __future__ import annotations

import io
from datetime import date, datetime, timedelta
from pathlib import Path
from xml.sax.saxutils import escape

from app.models.trip import PackingList, TripPlan
from app.services import packing_service

# ==================== 共用内容模型 ====================
# 下面这些函数产出的是"中立结构"（元组 / 字典 / 列表），
# 不含任何 docx 或 reportlab 的类型 —— 保证两个渲染器看到的数据完全一致。


def date_range(plan: TripPlan) -> str:
    """行程日期区间，如 2026-10-16 ~ 2026-10-18。"""
    start = plan.start_date or date.today()
    end = start + timedelta(days=max(plan.days - 1, 0))
    return f"{start.isoformat()} ~ {end.isoformat()}"


def generator_label(plan: TripPlan) -> str:
    """编排来源的可读标签 —— 导出的文档里也要如实标注。"""
    if plan.generator == "llm":
        base = "AI 编排（大模型从候选池选点）"
    else:
        base = "规则生成"
    return f"{base} · {plan.generator_note}" if plan.generator_note else base


def total_budget(plan: TripPlan) -> float:
    return round(sum(b.amount for b in plan.budget), 2)


def summary_rows(plan: TripPlan) -> list[tuple[str, str]]:
    """文档开头的概览表：行程单的"抬头"。"""
    rows = [
        ("行程城市", " → ".join(plan.cities)),
        ("总天数", f"{plan.days} 天"),
        ("日期", date_range(plan)),
        ("节奏", f"{plan.pace}节奏"),
        ("预算（人均估算）", f"¥{total_budget(plan):,.2f}"),
        ("编排来源", generator_label(plan)),
    ]
    sources = sorted({a.source for d in plan.itinerary for a in d.activities if a.source})
    if sources:
        rows.append(("素材出处", " / ".join(sources)))
    return rows


def day_blocks(plan: TripPlan) -> list[dict]:
    """逐日行程：每天一个块，含活动表、餐食、住宿。"""
    start = plan.start_date or date.today()
    blocks: list[dict] = []
    for d in plan.itinerary:
        d0 = start + timedelta(days=max(d.day - 1, 0))
        blocks.append({
            "heading": f"D{d.day} · {d.city}" + (f" · {d.theme}" if d.theme else ""),
            "date": d0.isoformat(),
            "city": d.city,
            "theme": d.theme,
            # 活动行：(类型, 名称, 备注, 出处)
            "activities": [
                (a.kind, a.name, a.note or "", a.source or "—") for a in d.activities
            ] or [("—", "自由安排", "", "—")],
            "meals": list(d.meals),
            "lodging": d.lodging or "",
        })
    return blocks


def budget_rows(plan: TripPlan) -> list[tuple[str, str]]:
    """预算拆分行 + 合计，金额统一带 ¥ 与千分位。"""
    rows = [(b.label, f"¥{b.amount:,.2f}") for b in plan.budget]
    rows.append(("合计（人均估算）", f"¥{total_budget(plan):,.2f}"))
    return rows


def packing_blocks(packing: PackingList) -> list[tuple[str, list[str]]]:
    return [(c.category, list(c.items)) for c in packing.categories]


DISCLAIMER = (
    "预算是按城市消费水平、天数与景点数量估算的人均值，不含往返大交通与购物，"
    "实际花费随季节与个人选择浮动。行程内容来自本地攻略库与高德 POI 候选池，"
    "大模型只在候选池范围内选点，池外地名已被丢弃。"
)


def export_filename(plan: TripPlan, ext: str) -> str:
    """给下载文件起名，如 xingji_西安_2026-10-16.docx。

    城市名一律过滤掉路径分隔符 —— 文件名是拼出来的，
    不能假设里面永远没有奇怪的字符。
    """
    cities = "_".join(c.replace("/", "").replace("\\", "") for c in plan.cities[:3])
    d = (plan.start_date or date.today()).isoformat()
    return f"xingji_{cities}_{d}.{ext.lstrip('.')}"


# ==================== 渲染器一：Word（python-docx） ====================

_DOCX_FONT = "微软雅黑"


def _docx_apply_font(doc) -> None:
    """把 Normal 样式的中西文字体都设成微软雅黑。

    【为什么必须单独设 eastAsia】
        python-docx 的 font.name 只写 w:ascii / w:hAnsi 两个属性，
        中文字符走的是 w:eastAsia —— 漏掉它，正文中文会落回默认宋体，
        标题看起来就会和正文不是一套字体。
    """
    from docx.oxml.ns import qn

    style = doc.styles["Normal"]
    style.font.name = _DOCX_FONT
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), _DOCX_FONT)
    rfonts.set(qn("w:ascii"), _DOCX_FONT)
    rfonts.set(qn("w:hAnsi"), _DOCX_FONT)


def _docx_set_cell(cell, text: str, *, bold: bool = False) -> None:
    cell.text = text
    for p in cell.paragraphs:
        for run in p.runs:
            run.bold = bold
            run.font.name = _DOCX_FONT
            from docx.oxml.ns import qn
            run._element.rPr.rFonts.set(qn("w:eastAsia"), _DOCX_FONT)


def build_docx(plan: TripPlan, packing: PackingList | None = None) -> bytes:
    """把行程渲染成 .docx 字节流。"""
    from docx import Document
    from docx.shared import Pt, RGBColor

    packing = packing or packing_service.build_packing_list(plan)
    doc = Document()
    _docx_apply_font(doc)

    # ---- 抬头 ----
    doc.add_heading(f"{' → '.join(plan.cities)} · {plan.days} 天行程单", level=0)

    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, value in summary_rows(plan):
        row = table.add_row().cells
        _docx_set_cell(row[0], label, bold=True)
        _docx_set_cell(row[1], value)

    # ---- 逐日行程 ----
    doc.add_heading("逐日行程", level=1)
    for block in day_blocks(plan):
        doc.add_heading(f"{block['heading']}（{block['date']}）", level=2)

        t = doc.add_table(rows=1, cols=4)
        t.style = "Light Grid Accent 1"
        for cell, head in zip(t.rows[0].cells, ("类型", "安排", "备注", "出处")):
            _docx_set_cell(cell, head, bold=True)
        for kind, name, note, source in block["activities"]:
            cells = t.add_row().cells
            for cell, text in zip(cells, (kind, name, note, source)):
                _docx_set_cell(cell, text)

        if block["meals"]:
            p = doc.add_paragraph()
            p.add_run("推荐餐食：").bold = True
            p.add_run("、".join(block["meals"]))
        if block["lodging"]:
            p = doc.add_paragraph()
            p.add_run("住宿区域：").bold = True
            p.add_run(block["lodging"])

    # ---- 预算 ----
    doc.add_heading("预算拆分", level=1)
    bt = doc.add_table(rows=0, cols=2)
    bt.style = "Light Grid Accent 1"
    for label, amount in budget_rows(plan):
        cells = bt.add_row().cells
        # 合计行加粗，让"一眼看到总数"这件事在纸上也成立
        is_total = label.startswith("合计")
        _docx_set_cell(cells[0], label, bold=is_total)
        _docx_set_cell(cells[1], amount, bold=is_total)

    # ---- 打包清单 ----
    doc.add_heading("打包清单", level=1)
    for category, items in packing_blocks(packing):
        doc.add_heading(category, level=3)
        for item in items:
            doc.add_paragraph(item, style="List Bullet")

    # ---- 说明 ----
    doc.add_heading("说明", level=1)
    note = doc.add_paragraph(DISCLAIMER)
    note.runs[0].font.size = Pt(9)
    note.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    stamp = doc.add_paragraph(
        f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')} ｜ 行迹 · 智能行程规划"
    )
    stamp.runs[0].font.size = Pt(9)
    stamp.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ==================== 渲染器二：PDF（reportlab） ====================

# 优先用系统字体（字形好看、粗体真实），失败则退回内置 CID 字体。
# STSong-Light 是 reportlab 自带的，任何机器上都能出中文 —— 这是保底。
_PDF_FONT_CACHE: tuple[str, str] | None = None

_PDF_TTF_CANDIDATES = (
    (r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\msyhbd.ttc"),
    (r"C:\Windows\Fonts\simhei.ttf", r"C:\Windows\Fonts\simhei.ttf"),
    (r"C:\Windows\Fonts\simsun.ttc", r"C:\Windows\Fonts\simsun.ttc"),
)


def _pdf_fonts() -> tuple[str, str]:
    """返回 (正文字体名, 粗体字体名)，结果缓存，避免每次导出重复注册。"""
    global _PDF_FONT_CACHE
    if _PDF_FONT_CACHE:
        return _PDF_FONT_CACHE

    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont

    for regular_path, bold_path in _PDF_TTF_CANDIDATES:
        if not (Path(regular_path).exists() and Path(bold_path).exists()):
            continue
        try:
            pdfmetrics.registerFont(TTFont("XingjiRegular", regular_path, subfontIndex=0))
            pdfmetrics.registerFont(TTFont("XingjiBold", bold_path, subfontIndex=0))
            _PDF_FONT_CACHE = ("XingjiRegular", "XingjiBold")
            return _PDF_FONT_CACHE
        except Exception:
            continue   # 这个字体坏了就试下一个，不要因此让导出整个失败

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    _PDF_FONT_CACHE = ("STSong-Light", "STSong-Light")
    return _PDF_FONT_CACHE


def _p(text: str, style):
    """把文本包成 Paragraph，转义 XML 特殊字符。

    活动名里出现 & 或 < 不罕见（"牛肉面 & 小吃"），
    不转义会让 reportlab 直接抛 XML 解析错误，整个导出失败。
    """
    from reportlab.platypus import Paragraph

    return Paragraph(escape(str(text)), style)


def build_pdf(plan: TripPlan, packing: PackingList | None = None) -> bytes:
    """把行程渲染成 PDF 字节流。"""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        KeepTogether, PageBreak, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    packing = packing or packing_service.build_packing_list(plan)
    regular, bold = _pdf_fonts()

    title = ParagraphStyle("t", fontName=bold, fontSize=20, leading=26,
                           textColor=colors.HexColor("#1f6f5c"), spaceAfter=4)
    h1 = ParagraphStyle("h1", fontName=bold, fontSize=14, leading=19,
                        textColor=colors.HexColor("#1f6f5c"), spaceBefore=12, spaceAfter=6)
    h2 = ParagraphStyle("h2", fontName=bold, fontSize=11.5, leading=16, spaceBefore=8, spaceAfter=4)
    body = ParagraphStyle("b", fontName=regular, fontSize=9.5, leading=14, alignment=TA_LEFT)
    cell = ParagraphStyle("c", parent=body, fontSize=9, leading=12.5)
    cell_head = ParagraphStyle("ch", parent=cell, fontName=bold)
    small = ParagraphStyle("s", parent=body, fontSize=8, leading=11,
                           textColor=colors.HexColor("#777777"))

    grid = TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d5d5d5")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef5f2")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ])

    story: list = [_p(f"{' → '.join(plan.cities)} · {plan.days} 天行程单", title)]

    # ---- 概览 ----
    overview = Table(
        [[_p(k, cell_head), _p(v, cell)] for k, v in summary_rows(plan)],
        colWidths=[38 * mm, 130 * mm],
    )
    overview.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d5d5d5")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f6faf8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
    ]))
    story += [overview, _p("逐日行程", h1)]

    # ---- 逐日行程 ----
    for block in day_blocks(plan):
        rows = [[_p(h, cell_head) for h in ("类型", "安排", "备注", "出处")]]
        rows += [[_p(x, cell) for x in act] for act in block["activities"]]
        t = Table(rows, colWidths=[16 * mm, 46 * mm, 62 * mm, 44 * mm], repeatRows=1)
        t.setStyle(grid)

        head = _p(f"{block['heading']}（{block['date']}）", h2)
        tail: list = []
        if block["meals"]:
            tail.append(_p("推荐餐食：" + "、".join(block["meals"]), body))
        if block["lodging"]:
            tail.append(_p("住宿区域：" + block["lodging"], body))

        # 一天的内容尽量不跨页：标题和数据表被拆开读起来会错位
        story.append(KeepTogether([head, t] + tail))

    # ---- 预算 ----
    story.append(PageBreak())
    story.append(_p("预算拆分", h1))
    rows = [[_p("项目", cell_head), _p("金额", cell_head)]]
    for label, amount in budget_rows(plan):
        rows.append([_p(label, cell), _p(amount, cell)])
    bt = Table(rows, colWidths=[60 * mm, 40 * mm], repeatRows=1)
    bt.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d5d5d5")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef5f2")),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#f6faf8")),
        ("FONTNAME", (0, -1), (-1, -1), bold),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(bt)

    # ---- 打包清单 ----
    story.append(_p("打包清单", h1))
    for category, items in packing_blocks(packing):
        story.append(_p(category, h2))
        story.append(_p("　".join(f"□ {i}" for i in items), body))

    # ---- 说明 ----
    story += [
        Spacer(1, 8 * mm),
        _p("说明", h2),
        _p(DISCLAIMER, small),
        _p(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')} ｜ 行迹 · 智能行程规划", small),
    ]

    buf = io.BytesIO()
    SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
        title=f"{'、'.join(plan.cities)} 行程单", author="行迹 · 智能行程规划",
    ).build(story)
    return buf.getvalue()
